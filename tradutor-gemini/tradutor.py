"""
🌐 TRADUTOR GEMINI - Multi Idiomas
Traduz roteiros para vários idiomas usando Google Gemini AI

Autor: Claude
Data: 2024
"""

from flask import Flask, render_template_string, request, jsonify, send_file
import google.generativeai as genai
from docx import Document
from docx.shared import Pt
from io import BytesIO
import zipfile
import re

app = Flask(__name__)

# ============================================
# HTML COMPLETO (EMBUTIDO)
# ============================================

HTML_TEMPLATE = """
<!DOCTYPE html>
<html lang="pt-BR">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Tradutor Gemini</title>
    <style>
        * {
            margin: 0;
            padding: 0;
            box-sizing: border-box;
        }

        :root {
            --primary: #6366f1;
            --primary-dark: #4f46e5;
            --secondary: #10b981;
            --text: #1e293b;
            --text-light: #64748b;
            --border: #e2e8f0;
            --success: #10b981;
            --error: #ef4444;
        }

        body {
            font-family: 'Segoe UI', system-ui, sans-serif;
            background: linear-gradient(135deg, #667eea 0%, #764ba2 100%);
            min-height: 100vh;
            padding: 20px;
            color: var(--text);
        }

        .container {
            max-width: 900px;
            margin: 0 auto;
        }

        .main-header {
            text-align: center;
            margin-bottom: 30px;
            color: white;
        }

        .main-header h1 {
            font-size: 3rem;
            font-weight: 800;
            margin-bottom: 10px;
            text-shadow: 0 4px 6px rgba(0, 0, 0, 0.1);
        }

        .main-header p {
            font-size: 1.1rem;
            opacity: 0.95;
        }

        .main-card {
            background: white;
            border-radius: 24px;
            padding: 40px;
            box-shadow: 0 20px 60px rgba(0, 0, 0, 0.3);
            margin-bottom: 20px;
        }

        .input-section {
            margin-bottom: 35px;
        }

        .section-header {
            display: flex;
            justify-content: space-between;
            align-items: center;
            margin-bottom: 20px;
        }

        .section-header h3 {
            font-size: 1.3rem;
            font-weight: 700;
        }

        .modern-input {
            width: 100%;
            padding: 16px 20px;
            border: 2px solid var(--border);
            border-radius: 12px;
            font-size: 1rem;
            transition: all 0.3s;
            font-family: 'Courier New', monospace;
        }

        .modern-input:focus {
            outline: none;
            border-color: var(--primary);
            box-shadow: 0 0 0 4px rgba(99, 102, 241, 0.1);
        }

        .help-link {
            display: inline-block;
            margin-top: 12px;
            color: var(--primary);
            text-decoration: none;
            font-weight: 600;
        }

        .select-all-btn {
            padding: 8px 20px;
            background: var(--primary);
            color: white;
            border: none;
            border-radius: 8px;
            font-weight: 600;
            cursor: pointer;
            transition: all 0.3s;
        }

        .select-all-btn:hover {
            background: var(--primary-dark);
            transform: scale(1.05);
        }

        .languages-container {
            display: grid;
            grid-template-columns: repeat(auto-fill, minmax(180px, 1fr));
            gap: 12px;
        }

        .lang-item {
            display: flex;
            align-items: center;
            padding: 14px 18px;
            border: 2px solid var(--border);
            border-radius: 12px;
            cursor: pointer;
            transition: all 0.3s;
            background: #fafafa;
        }

        .lang-item:hover {
            border-color: var(--primary);
            background: rgba(99, 102, 241, 0.05);
            transform: translateY(-2px);
        }

        .lang-item input {
            width: 20px;
            height: 20px;
            margin-right: 12px;
            cursor: pointer;
            accent-color: var(--primary);
        }

        .lang-item input:checked + .lang-label {
            color: var(--primary);
            font-weight: 700;
        }

        .modern-textarea {
            width: 100%;
            padding: 18px 20px;
            border: 2px solid var(--border);
            border-radius: 12px;
            font-size: 1rem;
            font-family: inherit;
            resize: vertical;
            transition: all 0.3s;
        }

        .modern-textarea:focus {
            outline: none;
            border-color: var(--primary);
            box-shadow: 0 0 0 4px rgba(99, 102, 241, 0.1);
        }

        .translate-btn {
            width: 100%;
            padding: 20px;
            background: linear-gradient(135deg, var(--primary), var(--primary-dark));
            color: white;
            border: none;
            border-radius: 12px;
            font-size: 1.3rem;
            font-weight: 700;
            cursor: pointer;
            transition: all 0.3s;
            box-shadow: 0 10px 30px rgba(99, 102, 241, 0.4);
        }

        .translate-btn:hover:not(:disabled) {
            transform: translateY(-3px);
            box-shadow: 0 15px 40px rgba(99, 102, 241, 0.5);
        }

        .translate-btn:disabled {
            opacity: 0.6;
            cursor: not-allowed;
        }

        .progress-card, .results-card {
            background: white;
            border-radius: 24px;
            padding: 40px;
            box-shadow: 0 20px 60px rgba(0, 0, 0, 0.3);
            margin-bottom: 20px;
            text-align: center;
        }

        .progress-bar-container {
            position: relative;
            width: 100%;
            height: 40px;
            background: #e2e8f0;
            border-radius: 20px;
            overflow: hidden;
            margin: 20px 0;
        }

        .progress-bar-fill {
            height: 100%;
            background: linear-gradient(90deg, var(--primary), var(--secondary));
            width: 0%;
            transition: width 0.5s ease;
        }

        .progress-percent {
            position: absolute;
            top: 50%;
            left: 50%;
            transform: translate(-50%, -50%);
            font-weight: 700;
        }

        .download-all-box {
            background: linear-gradient(135deg, rgba(99, 102, 241, 0.1), rgba(16, 185, 129, 0.1));
            padding: 25px;
            border-radius: 16px;
            margin-bottom: 30px;
        }

        .download-all-btns {
            display: flex;
            gap: 12px;
            justify-content: center;
            flex-wrap: wrap;
            margin-top: 15px;
        }

        .dl-btn {
            padding: 12px 30px;
            background: white;
            border: 2px solid var(--primary);
            color: var(--primary);
            border-radius: 10px;
            font-weight: 700;
            cursor: pointer;
            transition: all 0.3s;
        }

        .dl-btn:hover {
            background: var(--primary);
            color: white;
            transform: translateY(-2px);
        }

        .translations-list {
            display: flex;
            flex-direction: column;
            gap: 20px;
        }

        .translation-item {
            border: 2px solid var(--border);
            border-radius: 16px;
            padding: 25px;
            background: #fafafa;
            transition: all 0.3s;
        }

        .translation-item:hover {
            border-color: var(--primary);
            transform: translateY(-2px);
            box-shadow: 0 5px 20px rgba(0, 0, 0, 0.1);
        }

        .translation-header {
            margin-bottom: 15px;
            padding-bottom: 15px;
            border-bottom: 2px solid var(--border);
        }

        .translation-title {
            font-size: 1.3rem;
            font-weight: 700;
        }

        .translation-preview {
            background: white;
            padding: 20px;
            border-radius: 12px;
            margin-bottom: 15px;
            font-size: 0.95rem;
            line-height: 1.6;
            max-height: 150px;
            overflow-y: auto;
            border: 1px solid var(--border);
        }

        .translation-actions {
            display: flex;
            gap: 10px;
            flex-wrap: wrap;
        }

        .btn-download {
            flex: 1;
            min-width: 100px;
            padding: 10px 20px;
            background: white;
            border: 2px solid var(--primary);
            color: var(--primary);
            border-radius: 8px;
            font-weight: 600;
            cursor: pointer;
            transition: all 0.3s;
        }

        .btn-download:hover {
            background: var(--primary);
            color: white;
            transform: translateY(-2px);
        }

        @media (max-width: 768px) {
            .main-header h1 { font-size: 2rem; }
            .main-card { padding: 25px; }
            .languages-container { grid-template-columns: 1fr; }
        }
    </style>
</head>
<body>
    <div class="container">
        <header class="main-header">
            <h1>🌐 Tradutor Gemini</h1>
            <p>Traduza seus roteiros para vários idiomas instantaneamente</p>
        </header>

        <div class="main-card">
            <div class="input-section">
                <div class="section-header">
                    <h3>🔑 Configuração</h3>
                </div>
                <input type="password" id="apiKey" placeholder="Cole sua API Key do Gemini aqui" class="modern-input">
                <a href="https://makersuite.google.com/app/apikey" target="_blank" class="help-link">
                    ➜ Não tem API Key? Clique aqui para obter
                </a>
            </div>

            <div class="input-section">
                <div class="section-header">
                    <h3>🗣️ Escolha os Idiomas</h3>
                    <button id="selectAll" class="select-all-btn">Todos</button>
                </div>
                <div class="languages-container">
                    <label class="lang-item">
                        <input type="checkbox" value="es" data-lang="Espanhol">
                        <span class="lang-label">🇪🇸 Espanhol</span>
                    </label>
                    <label class="lang-item">
                        <input type="checkbox" value="en" data-lang="Inglês">
                        <span class="lang-label">🇺🇸 Inglês</span>
                    </label>
                    <label class="lang-item">
                        <input type="checkbox" value="ru" data-lang="Russo">
                        <span class="lang-label">🇷🇺 Russo</span>
                    </label>
                    <label class="lang-item">
                        <input type="checkbox" value="ar" data-lang="Árabe">
                        <span class="lang-label">🇸🇦 Árabe</span>
                    </label>
                    <label class="lang-item">
                        <input type="checkbox" value="tr" data-lang="Turco">
                        <span class="lang-label">🇹🇷 Turco</span>
                    </label>
                    <label class="lang-item">
                        <input type="checkbox" value="pl" data-lang="Polonês">
                        <span class="lang-label">🇵🇱 Polonês</span>
                    </label>
                    <label class="lang-item">
                        <input type="checkbox" value="de" data-lang="Alemão">
                        <span class="lang-label">🇩🇪 Alemão</span>
                    </label>
                    <label class="lang-item">
                        <input type="checkbox" value="fr" data-lang="Francês">
                        <span class="lang-label">🇫🇷 Francês</span>
                    </label>
                    <label class="lang-item">
                        <input type="checkbox" value="it" data-lang="Italiano">
                        <span class="lang-label">🇮🇹 Italiano</span>
                    </label>
                    <label class="lang-item">
                        <input type="checkbox" value="ro" data-lang="Romeno">
                        <span class="lang-label">🇷🇴 Romeno</span>
                    </label>
                </div>
            </div>

            <div class="input-section">
                <div class="section-header">
                    <h3>📝 Seu Roteiro em Português</h3>
                </div>
                <textarea id="scriptText" placeholder="Cole seu roteiro aqui..." class="modern-textarea" rows="10"></textarea>
            </div>

            <button id="translateBtn" class="translate-btn">🚀 Traduzir Agora</button>
        </div>

        <div id="progressSection" class="progress-card" style="display: none;">
            <h3>⏳ Traduzindo...</h3>
            <div class="progress-bar-container">
                <div id="progressFill" class="progress-bar-fill"></div>
                <span id="progressPercent" class="progress-percent">0%</span>
            </div>
            <p id="progressText">Iniciando...</p>
        </div>

        <div id="resultsSection" class="results-card" style="display: none;">
            <h3>✅ Traduções Prontas!</h3>
            <div class="download-all-box">
                <p><strong>Baixar todos os idiomas:</strong></p>
                <div class="download-all-btns">
                    <button class="dl-btn" onclick="downloadAll('txt')">📄 TXT</button>
                    <button class="dl-btn" onclick="downloadAll('docx')">📘 DOCX</button>
                    <button class="dl-btn" onclick="downloadAll('srt')">🎬 SRT</button>
                </div>
            </div>
            <div id="translationResults" class="translations-list"></div>
        </div>
    </div>

    <script>
        let translations = {};
        const apiKeyInput = document.getElementById('apiKey');
        const translateBtn = document.getElementById('translateBtn');
        const scriptTextArea = document.getElementById('scriptText');
        const progressSection = document.getElementById('progressSection');
        const resultsSection = document.getElementById('resultsSection');
        const progressFill = document.getElementById('progressFill');
        const progressPercent = document.getElementById('progressPercent');
        const progressText = document.getElementById('progressText');
        const translationResults = document.getElementById('translationResults');

        // Carregar API Key salva
        window.addEventListener('load', () => {
            const savedKey = localStorage.getItem('gemini_api_key');
            if (savedKey) apiKeyInput.value = savedKey;
        });

        apiKeyInput.addEventListener('input', () => {
            localStorage.setItem('gemini_api_key', apiKeyInput.value);
        });

        // Selecionar todos
        document.getElementById('selectAll').addEventListener('click', () => {
            const checkboxes = document.querySelectorAll('.lang-item input[type="checkbox"]');
            const allChecked = Array.from(checkboxes).every(cb => cb.checked);
            checkboxes.forEach(cb => cb.checked = !allChecked);
        });

        // Traduzir
        translateBtn.addEventListener('click', async () => {
            const apiKey = apiKeyInput.value.trim();
            const text = scriptTextArea.value.trim();
            const languages = Array.from(document.querySelectorAll('.lang-item input:checked'))
                .map(cb => ({ code: cb.value, name: cb.getAttribute('data-lang') }));

            if (!apiKey) {
                alert('❌ Por favor, cole sua API Key do Gemini!');
                return;
            }

            if (languages.length === 0) {
                alert('❌ Por favor, selecione pelo menos um idioma!');
                return;
            }

            if (!text) {
                alert('❌ Por favor, cole seu roteiro!');
                return;
            }

            await startTranslation(apiKey, languages, text);
        });

        async function startTranslation(apiKey, languages, text) {
            translations = {};
            resultsSection.style.display = 'none';
            progressSection.style.display = 'block';
            translateBtn.disabled = true;

            let completed = 0;
            const total = languages.length;

            for (const language of languages) {
                progressText.textContent = `Traduzindo para ${language.name}...`;

                try {
                    const response = await fetch('/translate', {
                        method: 'POST',
                        headers: { 'Content-Type': 'application/json' },
                        body: JSON.stringify({ apiKey, text, targetLanguage: language.code, languageName: language.name })
                    });

                    if (!response.ok) {
                        const errorData = await response.json();
                        throw new Error(errorData.error || 'Erro desconhecido');
                    }

                    const result = await response.json();
                    translations[language.code] = { name: language.name, text: result.translation };
                } catch (error) {
                    console.error(`Erro em ${language.name}:`, error);
                    translations[language.code] = { name: language.name, text: `❌ Erro: ${error.message}`, error: true };
                }

                completed++;
                const percentage = Math.round((completed / total) * 100);
                progressFill.style.width = `${percentage}%`;
                progressPercent.textContent = `${percentage}%`;
            }

            progressText.textContent = '✅ Traduções concluídas!';
            translateBtn.disabled = false;

            setTimeout(() => {
                progressSection.style.display = 'none';
                displayResults();
            }, 1500);
        }

        function displayResults() {
            translationResults.innerHTML = '';

            Object.entries(translations).forEach(([code, data]) => {
                const item = document.createElement('div');
                item.className = 'translation-item';

                const preview = data.error
                    ? `<p style="color: #ef4444; font-weight: 600;">${data.text}</p>`
                    : data.text.substring(0, 250) + (data.text.length > 250 ? '...' : '');

                item.innerHTML = `
                    <div class="translation-header">
                        <h4 class="translation-title">${data.name}</h4>
                    </div>
                    <div class="translation-preview">${preview}</div>
                    ${!data.error ? `
                    <div class="translation-actions">
                        <button class="btn-download" onclick="downloadSingle('${code}', 'txt')">📄 TXT</button>
                        <button class="btn-download" onclick="downloadSingle('${code}', 'docx')">📘 DOCX</button>
                        <button class="btn-download" onclick="downloadSingle('${code}', 'srt')">🎬 SRT</button>
                    </div>
                    ` : ''}
                `;

                translationResults.appendChild(item);
            });

            resultsSection.style.display = 'block';
            resultsSection.scrollIntoView({ behavior: 'smooth' });
        }

        async function downloadSingle(languageCode, format) {
            const data = translations[languageCode];
            if (!data || data.error) return;

            try {
                const response = await fetch('/generate-file', {
                    method: 'POST',
                    headers: { 'Content-Type': 'application/json' },
                    body: JSON.stringify({ text: data.text, language: data.name, format })
                });

                if (!response.ok) throw new Error('Erro ao gerar arquivo');

                const blob = await response.blob();
                const url = window.URL.createObjectURL(blob);
                const a = document.createElement('a');
                a.href = url;
                a.download = `${data.name.toLowerCase().replace(/\\s+/g, '_')}.${format}`;
                a.click();
                window.URL.revokeObjectURL(url);
            } catch (error) {
                alert('❌ Erro ao baixar: ' + error.message);
            }
        }

        async function downloadAll(format) {
            const validTranslations = Object.entries(translations)
                .filter(([_, data]) => !data.error)
                .map(([code, data]) => ({ language: data.name, text: data.text }));

            if (validTranslations.length === 0) {
                alert('❌ Nenhuma tradução válida');
                return;
            }

            try {
                const response = await fetch('/generate-all', {
                    method: 'POST',
                    headers: { 'Content-Type': 'application/json' },
                    body: JSON.stringify({ translations: validTranslations, format })
                });

                if (!response.ok) throw new Error('Erro ao gerar ZIP');

                const blob = await response.blob();
                const url = window.URL.createObjectURL(blob);
                const a = document.createElement('a');
                a.href = url;
                a.download = `todas_traducoes_${format}.zip`;
                a.click();
                window.URL.revokeObjectURL(url);
            } catch (error) {
                alert('❌ Erro ao baixar: ' + error.message);
            }
        }
    </script>
</body>
</html>
"""

# ============================================
# MAPEAMENTO DE IDIOMAS
# ============================================

LANGUAGE_NAMES = {
    'es': 'Spanish',
    'en': 'English',
    'ru': 'Russian',
    'ar': 'Arabic',
    'tr': 'Turkish',
    'pl': 'Polish',
    'de': 'German',
    'fr': 'French',
    'it': 'Italian',
    'ro': 'Romanian'
}

# ============================================
# FUNÇÕES AUXILIARES
# ============================================

def split_text_for_srt(text, max_chars=400):
    """Divide texto em chunks de 400 caracteres para SRT"""
    sentences = re.findall(r'[^.!?]+[.!?]+', text) or [text]
    chunks = []
    current_chunk = ''

    for sentence in sentences:
        sentence = sentence.strip()

        if len(current_chunk + ' ' + sentence) <= max_chars:
            current_chunk += (' ' if current_chunk else '') + sentence
        else:
            if current_chunk:
                chunks.append(current_chunk)

            if len(sentence) > max_chars:
                words = sentence.split()
                word_chunk = ''
                for word in words:
                    if len(word_chunk + ' ' + word) <= max_chars:
                        word_chunk += (' ' if word_chunk else '') + word
                    else:
                        if word_chunk:
                            chunks.append(word_chunk)
                        word_chunk = word
                if word_chunk:
                    current_chunk = word_chunk
                else:
                    current_chunk = ''
            else:
                current_chunk = sentence

    if current_chunk:
        chunks.append(current_chunk)

    return chunks

def format_srt_time(seconds):
    """Formata tempo para SRT (HH:MM:SS,mmm)"""
    hours = int(seconds // 3600)
    minutes = int((seconds % 3600) // 60)
    secs = int(seconds % 60)
    milliseconds = int((seconds % 1) * 1000)
    return f"{hours:02d}:{minutes:02d}:{secs:02d},{milliseconds:03d}"

def generate_srt(text):
    """Gera conteúdo SRT"""
    chunks = split_text_for_srt(text)
    srt_content = ''

    for index, chunk in enumerate(chunks):
        start_time = format_srt_time(index * 3)
        end_time = format_srt_time((index + 1) * 3)
        srt_content += f"{index + 1}\n{start_time} --> {end_time}\n{chunk}\n\n"

    return srt_content

def generate_docx(text, language):
    """Gera arquivo DOCX"""
    doc = Document()

    # Título
    title = doc.add_heading(f'Tradução - {language}', 0)
    title.alignment = 1  # Centralizar

    # Conteúdo
    paragraphs = text.split('\n')
    for para in paragraphs:
        if para.strip():
            p = doc.add_paragraph(para)
            p.style.font.size = Pt(12)

    # Salvar em memória
    file_stream = BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)
    return file_stream

# ============================================
# ROTAS
# ============================================

@app.route('/')
def index():
    """Página principal"""
    return render_template_string(HTML_TEMPLATE)

@app.route('/translate', methods=['POST'])
def translate():
    """Endpoint de tradução"""
    try:
        data = request.json
        api_key = data.get('apiKey')
        text = data.get('text')
        target_language = data.get('targetLanguage')
        language_name = data.get('languageName')

        if not api_key or not text or not target_language:
            return jsonify({'error': 'Parâmetros inválidos'}), 400

        # Configurar Gemini
        genai.configure(api_key=api_key)
        model = genai.GenerativeModel('gemini-1.5-pro-latest')

        # Criar prompt
        full_language_name = LANGUAGE_NAMES.get(target_language, language_name)
        prompt = f"""Translate the following Portuguese text to {full_language_name}.

IMPORTANT: Return ONLY the translation, without any explanations, notes, or additional text.

Text to translate:
{text}"""

        # Fazer tradução
        response = model.generate_content(prompt)
        translation = response.text.strip()

        return jsonify({'translation': translation})

    except Exception as e:
        print(f"Erro na tradução: {e}")
        return jsonify({'error': str(e)}), 500

@app.route('/generate-file', methods=['POST'])
def generate_file():
    """Gera arquivo individual"""
    try:
        data = request.json
        text = data.get('text')
        language = data.get('language')
        format_type = data.get('format')

        if not text or not format_type:
            return jsonify({'error': 'Parâmetros inválidos'}), 400

        filename = f"{language.lower().replace(' ', '_')}.{format_type}"

        if format_type == 'txt':
            file_stream = BytesIO(text.encode('utf-8'))
            mimetype = 'text/plain'

        elif format_type == 'docx':
            file_stream = generate_docx(text, language)
            mimetype = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'

        elif format_type == 'srt':
            srt_content = generate_srt(text)
            file_stream = BytesIO(srt_content.encode('utf-8'))
            mimetype = 'application/x-subrip'

        else:
            return jsonify({'error': 'Formato inválido'}), 400

        return send_file(
            file_stream,
            mimetype=mimetype,
            as_attachment=True,
            download_name=filename
        )

    except Exception as e:
        print(f"Erro ao gerar arquivo: {e}")
        return jsonify({'error': str(e)}), 500

@app.route('/generate-all', methods=['POST'])
def generate_all():
    """Gera ZIP com todos os arquivos"""
    try:
        data = request.json
        translations = data.get('translations', [])
        format_type = data.get('format')

        if not translations or not format_type:
            return jsonify({'error': 'Parâmetros inválidos'}), 400

        # Criar ZIP em memória
        zip_stream = BytesIO()

        with zipfile.ZipFile(zip_stream, 'w', zipfile.ZIP_DEFLATED) as zip_file:
            for translation in translations:
                language = translation['language']
                text = translation['text']
                filename = f"{language.lower().replace(' ', '_')}.{format_type}"

                if format_type == 'txt':
                    content = text.encode('utf-8')

                elif format_type == 'docx':
                    docx_stream = generate_docx(text, language)
                    content = docx_stream.read()

                elif format_type == 'srt':
                    srt_content = generate_srt(text)
                    content = srt_content.encode('utf-8')

                else:
                    continue

                zip_file.writestr(filename, content)

        zip_stream.seek(0)

        return send_file(
            zip_stream,
            mimetype='application/zip',
            as_attachment=True,
            download_name=f'todas_traducoes_{format_type}.zip'
        )

    except Exception as e:
        print(f"Erro ao gerar ZIP: {e}")
        return jsonify({'error': str(e)}), 500

# ============================================
# EXECUTAR APLICAÇÃO
# ============================================

if __name__ == '__main__':
    print("=" * 50)
    print("🌐 TRADUTOR GEMINI - Multi Idiomas")
    print("=" * 50)
    print("\n✅ Servidor iniciando...")
    print("📝 Acesse: http://localhost:5000")
    print("\n💡 Para parar: Ctrl+C")
    print("=" * 50)

    app.run(debug=True, host='0.0.0.0', port=5000)
